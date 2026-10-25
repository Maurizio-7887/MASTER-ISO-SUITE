"""Genera il Rapporto di Audit stampabile (.docx) nello stesso formato
storico dell'azienda (PG-AUD-02): tabella intestazione + tabella punti ISO
9001 con rilievi/azioni, a partire dai dati compilati online."""
import io
from flask import send_file
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _azienda():
    from flask import current_app
    return current_app.config.get("RAGIONE_SOCIALE", "Azienda")



def _bordi_cella(cella):
    tc = cella._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for lato in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lato}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _sfondo_cella(cella, colore_hex):
    tc = cella._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), colore_hex)
    tcPr.append(shd)


def genera_docx_rapporto_audit(report):
    doc = Document()
    sezione = doc.sections[0]
    sezione.left_margin = sezione.right_margin = Cm(1.5)
    sezione.top_margin = sezione.bottom_margin = Cm(1.5)

    # --- Intestazione: tabella a 3 colonne, come il modello originale ---
    intestazione = doc.add_table(rows=1, cols=3)
    intestazione.alignment = WD_TABLE_ALIGNMENT.CENTER
    larghezze = (Cm(3), Cm(11), Cm(3.5))
    for cella, larghezza in zip(intestazione.rows[0].cells, larghezze):
        cella.width = larghezza
        _bordi_cella(cella)
    intestazione.rows[0].cells[0].paragraphs[0].add_run(_azienda().upper()).bold = True
    p_titolo = intestazione.rows[0].cells[1].paragraphs[0]
    p_titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titolo = p_titolo.add_run("RAPPORTO DI AUDIT")
    run_titolo.bold = True
    run_titolo.font.size = Pt(16)
    p_modello = intestazione.rows[0].cells[2].paragraphs[0]
    p_modello.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_modello = p_modello.add_run("MODELLO SISTEMA QUALITA'")
    run_modello.bold = True

    doc.add_paragraph()

    # --- Paragrafo introduttivo ---
    intro = doc.add_paragraph()
    intro.add_run("Il gruppo di Audit, costituito da ")
    r = intro.add_run(report.team_lead or "_______________")
    r.bold = True; r.italic = True
    intro.add_run(" in qualità di responsabile del Gruppo di Audit")
    if report.team_members:
        intro.add_run(", ")
        r2 = intro.add_run(report.team_members)
        r2.italic = True
        intro.add_run(" in qualità di esperti d'area")
    intro.add_run(", hanno esaminato in data ")
    r3 = intro.add_run(report.audit_dates or "_______________")
    r3.bold = True
    intro.add_run(", il funzionamento dell'intero sistema qualità aziendale; i rilievi emersi sono riportati nel rapporto che segue:")

    doc.add_paragraph()

    # --- Tabella dei rilievi ---
    tabella = doc.add_table(rows=1, cols=4)
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    larghezze_col = (Cm(1.5), Cm(5), Cm(7), Cm(5))
    intestazioni = ("Sez.", "ISO 9001:2015", "Descrizione ANOMALIE VALUTAZIONE GIUDIZIO", "AZIONI DECISE/consigliate suggerite")
    for cella, testo, larghezza in zip(tabella.rows[0].cells, intestazioni, larghezze_col):
        cella.width = larghezza
        _bordi_cella(cella)
        _sfondo_cella(cella, "D9D9D9")
        p = cella.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(testo)
        run.bold = True
        run.font.size = Pt(9)

    for rilievo in report.findings:
        riga = tabella.add_row()
        celle = riga.cells
        for cella, larghezza in zip(celle, larghezze_col):
            cella.width = larghezza
            _bordi_cella(cella)
        if rilievo.is_heading:
            for cella in celle:
                _sfondo_cella(cella, "F2F2F2")
            celle[0].paragraphs[0].add_run(rilievo.sez).bold = True
            r = celle[1].paragraphs[0].add_run(rilievo.iso_title)
            r.bold = True
        else:
            celle[0].paragraphs[0].add_run(rilievo.sez)
            celle[1].paragraphs[0].add_run(rilievo.iso_title).bold = True
            celle[2].paragraphs[0].add_run(rilievo.findings_text or "")
            p_azioni = celle[3].paragraphs[0]
            p_azioni.add_run(rilievo.actions_text or "")
            if rilievo.attachments:
                p_all = celle[3].add_paragraph()
                run_all = p_all.add_run("ALLEGATI:")
                run_all.bold = True
                for allegato in rilievo.attachments:
                    celle[3].add_paragraph(f"- {allegato.label or (allegato.document.document_title if allegato.document else '')}")
        for cella in celle:
            for p in cella.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    # --- Piè di pagina ---
    for sezione_doc in doc.sections:
        piede = sezione_doc.footer.paragraphs[0]
        piede.text = f"{_azienda()} — File: Rapporto_Audit_{report.year}.docx — Stato: {report.status}"
        piede.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in piede.runs:
            run.font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    nome_file = f"Rapporto_Audit_{report.year}_{(report.audit_dates or '').replace(' ', '_')}.docx"
    return send_file(buffer, as_attachment=True, download_name=nome_file,
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
