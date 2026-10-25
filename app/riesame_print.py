"""Genera il Verbale di Riesame della Direzione stampabile (.docx) nello
stesso formato del modulo di riesame aziendale (MOD-RD), quello già noto al
certificatore, a partire dai dati compilati online (ManagementReview +
ManagementReviewActionItem)."""
import io
from flask import send_file
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _azienda():
    from flask import current_app
    return current_app.config.get("RAGIONE_SOCIALE", "Azienda")



def _bordi_cella(cella):
    tc = cella._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for lato in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lato}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _sfondo_cella(cella, colore_hex):
    tc = cella._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), colore_hex)
    tcPr.append(shd)


def _sezione(doc, titolo, testo):
    """Una sezione del verbale: titolo in grassetto su sfondo grigio (come
    nel formato storico a tabella singola), testo libero sotto."""
    tabella = doc.add_table(rows=2 if testo else 1, cols=1)
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    cella_titolo = tabella.rows[0].cells[0]
    _bordi_cella(cella_titolo)
    _sfondo_cella(cella_titolo, "D9D9D9")
    run = cella_titolo.paragraphs[0].add_run(titolo)
    run.bold = True
    run.font.size = Pt(10)
    if testo:
        cella_testo = tabella.rows[1].cells[0]
        _bordi_cella(cella_testo)
        for i, riga in enumerate(str(testo).split("\n")):
            p = cella_testo.paragraphs[0] if i == 0 else cella_testo.add_paragraph()
            r = p.add_run(riga)
            r.font.size = Pt(9)
    doc.add_paragraph()


def _sezione_analisi(doc, sezione, immagine_png=None):
    """Un punto del riesame nel formato analisi-per-punto richiesto
    dall'ispettore: titolo, 'Cosa va bene', 'Cosa non va / criticità', ed
    elenco degli allegati caricati per quel punto. Se disponibile, il
    grafico con i dati reali del programma viene inserito subito dopo."""
    tabella = doc.add_table(rows=1, cols=1)
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    cella_titolo = tabella.rows[0].cells[0]
    _bordi_cella(cella_titolo)
    _sfondo_cella(cella_titolo, "D9D9D9")
    run = cella_titolo.paragraphs[0].add_run(sezione.title)
    run.bold = True
    run.font.size = Pt(10)

    def _riga(etichetta, testo):
        riga = tabella.add_row()
        cella = riga.cells[0]
        _bordi_cella(cella)
        p = cella.paragraphs[0]
        p.add_run(f"{etichetta}: ").bold = True
        if testo:
            for i, linea in enumerate(str(testo).split("\n")):
                pp = p if i == 0 else cella.add_paragraph()
                pp.add_run(linea).font.size = Pt(9)
        for run_ in p.runs:
            run_.font.size = Pt(9)

    _riga("Cosa va bene", sezione.positives_text)
    _riga("Cosa non va / criticità", sezione.negatives_text)
    if sezione.attachments:
        riga = tabella.add_row()
        cella = riga.cells[0]
        _bordi_cella(cella)
        p = cella.paragraphs[0]
        p.add_run("Allegati: ").bold = True
        for allegato in sezione.attachments:
            p2 = cella.add_paragraph(f"- {allegato.label or (allegato.document.document_title if allegato.document else '')}")
            for run_ in p2.runs:
                run_.font.size = Pt(9)
    doc.add_paragraph()
    if immagine_png:
        doc.add_picture(io.BytesIO(immagine_png), width=Cm(14))
        doc.add_paragraph()


def _tabella_piani(doc, titolo, colonna_esito, items):
    tabella = doc.add_table(rows=1, cols=4)
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    larghezze = (Cm(0.9), Cm(6.5), Cm(3), Cm(3.5))
    intestazioni = ("N°", "ATTIVITÀ", "RISORSE / RESPONSABILE", colonna_esito)
    for cella, testo, larghezza in zip(tabella.rows[0].cells, intestazioni, larghezze):
        cella.width = larghezza
        _bordi_cella(cella)
        _sfondo_cella(cella, "D9D9D9")
        p = cella.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(testo)
        r.bold = True
        r.font.size = Pt(9)
    if not items:
        riga = tabella.add_row()
        for cella, larghezza in zip(riga.cells, larghezze):
            cella.width = larghezza
            _bordi_cella(cella)
        riga.cells[1].paragraphs[0].add_run("Nessuna voce.").italic = True
    for n, voce in enumerate(items, start=1):
        riga = tabella.add_row()
        celle = riga.cells
        for cella, larghezza in zip(celle, larghezze):
            cella.width = larghezza
            _bordi_cella(cella)
        celle[0].paragraphs[0].add_run(str(n))
        celle[1].paragraphs[0].add_run(voce.activity or "")
        risorse_resp = " / ".join(x for x in (voce.resources, voce.responsible) if x)
        celle[2].paragraphs[0].add_run(risorse_resp)
        celle[3].paragraphs[0].add_run(voce.status_text or "")
        for cella in celle:
            for p in cella.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def genera_docx_riesame_direzione(review, grafici_per_sezione=None):
    grafici_per_sezione = grafici_per_sezione or {}
    doc = Document()
    sezione_doc = doc.sections[0]
    sezione_doc.left_margin = sezione_doc.right_margin = Cm(1.5)
    sezione_doc.top_margin = sezione_doc.bottom_margin = Cm(1.5)

    # --- Intestazione ---
    intestazione = doc.add_table(rows=1, cols=3)
    intestazione.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cella, larghezza in zip(intestazione.rows[0].cells, (Cm(3), Cm(11), Cm(3.5))):
        cella.width = larghezza
        _bordi_cella(cella)
    intestazione.rows[0].cells[0].paragraphs[0].add_run(_azienda().upper()).bold = True
    p_titolo = intestazione.rows[0].cells[1].paragraphs[0]
    p_titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titolo = p_titolo.add_run("VERBALE RIESAME DELLA DIREZIONE")
    run_titolo.bold = True
    run_titolo.font.size = Pt(16)
    p_modello = intestazione.rows[0].cells[2].paragraphs[0]
    p_modello.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_modello.add_run("MOD-RD").bold = True

    doc.add_paragraph()

    # --- Paragrafo introduttivo ---
    intro = doc.add_paragraph()
    intro.add_run("La Direzione ha proceduto in data ")
    r = intro.add_run(review.review_date.strftime("%d/%m/%Y"))
    r.bold = True
    periodo = review.reference_period or str(review.review_date.year - 1)
    intro.add_run(f" al riesame del proprio Sistema Qualità relativo al {periodo}.")
    doc.add_paragraph("Il presente verbale ed i documenti da esso richiamati documentano l'esito del Riesame.")

    doc.add_paragraph()
    doc.add_paragraph("Il riesame si è basato sulle risultanze documentali di:")
    base_labels = (
        ("rapporti verifiche ispettive interne", review.basis_audit_reports),
        ("rapporti verifiche certificative", review.basis_certification_reports),
        ("azioni condotte in conseguenza del precedente riesame", review.basis_previous_review_actions),
        ("reclami ed informazioni di ritorno dai clienti", review.basis_customer_feedback),
        ("analisi andamento azioni correttive del SI", review.basis_corrective_actions_trend),
        ("analisi andamento azioni preventive del SI", review.basis_preventive_actions_trend),
        ("RAC avviate relativamente alle diverse aree/attività", review.basis_rac),
        ("analisi del contesto", review.basis_context_analysis),
    )
    for etichetta, valore in base_labels:
        doc.add_paragraph(f"{etichetta}: {'SI' if valore else 'NO'}.", style=None)

    doc.add_paragraph()

    sezioni_per_chiave = {s.section_key: s for s in review.sections}

    def _stampa_punto(chiave):
        s = sezioni_per_chiave.get(chiave)
        if s:
            _sezione_analisi(doc, s, immagine_png=grafici_per_sezione.get(chiave))

    _stampa_punto("risk_management")

    part_a = doc.add_paragraph()
    part_a.add_run("PARTE A: ELEMENTI IN USCITA").bold = True

    _stampa_punto("quality_policy")

    _tabella_piani(doc, "VERIFICHE AI PIANI DI MIGLIORAMENTO DELL'ANNO PRECEDENTE", "RISULTATI AL 31/12",
                    [i for i in review.action_items if i.kind == "verifica_precedente"])

    _stampa_punto("improvement_dev")
    _stampa_punto("internal_audit")
    _stampa_punto("certification_audit")
    _stampa_punto("complaints")
    _stampa_punto("nonconformities")
    _stampa_punto("injuries")
    _stampa_punto("training")
    _stampa_punto("suppliers")
    _stampa_punto("indicators")

    _tabella_piani(doc, "CONSIDERAZIONI E PROPOSTE DI MIGLIORAMENTO", "DATA EFFETTUAZIONE",
                    [i for i in review.action_items if i.kind == "nuova_proposta"])

    _stampa_punto("previous_review")

    _sezione(doc, "CONCLUSIONI", review.conclusions)

    doc.add_paragraph()
    p_firme = doc.add_paragraph()
    p_firme.add_run("I partecipanti al Riesame (Firma): ")
    p_firme.add_run(review.participants or "_" * 60)

    p_visto = doc.add_paragraph()
    p_visto.add_run("Visto: Direzione ______________________  ")
    data_visto = review.signature_date.strftime("%d/%m/%y") if review.signature_date else "__________"
    p_visto.add_run(f"Data {data_visto}")

    if review.notes:
        doc.add_paragraph()
        p_note = doc.add_paragraph()
        p_note.add_run("Note: ").bold = True
        p_note.add_run(review.notes)

    for sez in doc.sections:
        piede = sez.footer.paragraphs[0]
        piede.text = f"MOD-RD — Verbale Riesame della Direzione del {review.review_date.strftime('%d/%m/%Y')}"
        piede.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in piede.runs:
            run.font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    nome_file = f"Riesame_Direzione_{review.review_date.strftime('%Y%m%d')}.docx"
    return send_file(buffer, as_attachment=True, download_name=nome_file,
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
